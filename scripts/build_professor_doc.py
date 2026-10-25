"""Guida Word aggiornata da sorgente verificata. Eseguire con Python bundled."""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]


def add_runs(p,text):
    for i,chunk in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        p.add_run(chunk).bold=bool(i%2)


def table(document,stats):
    t=document.add_table(rows=1,cols=4)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[1.4,1.45,1.65,1.3]
    headers=['Clip','Etichette','Annotabili','Frame']
    for i,(c,h) in enumerate(zip(t.rows[0].cells,headers)):
        c.text=h;c.width=Inches(widths[i])
        pr=c._tc.get_or_add_tcPr();fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EDF2');pr.append(fill)
        for run in c.paragraphs[0].runs:run.bold=True
    for record in stats['per_clip']:
        row=t.add_row()
        for i,v in enumerate([record['clip'],record['labels'],record['annotatable'],record['frames']]):
            row.cells[i].text=str(v);row.cells[i].width=Inches(widths[i])
    total=t.add_row()
    for c,v in zip(total.cells,['Totale',stats['labels'],stats['annotatable'],stats['frames']]):
        c.text=str(v)
        for run in c.paragraphs[0].runs:run.bold=True
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    t._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in t.rows:
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');margins.append(x)
            c._tc.get_or_add_tcPr().append(margins)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.
                for run in p.runs:run.font.size=Pt(10.5)
    document.add_paragraph().paragraph_format.space_after=Pt(2)


def main():
    stats=json.loads((ROOT/'reports/annotazioni_complete_20261020/status.json').read_text(encoding='utf-8'))
    source=(ROOT/'reports/GUIDA_INCONTRO_PROFESSORE_2026-10-20.md').read_text(encoding='utf-8')
    source=source.replace('Sto preparando la verifica quantitativa\nsu sette clip separate, con annotazioni manuali.',
                          'Ho completato le annotazioni manuali\nsu sette clip separate e sto preparando la verifica quantitativa.')
    source=source.replace('- Annotazioni presenti: 104 etichette, 84 annotabili e 20 non annotabili, su 16 frame.',
                          '- Annotazioni presenti: 216 etichette, 176 annotabili e 40 non annotabili, su 40 frame.')
    source=source.replace('- Distribuzione: clip_01 64 etichette, clip_03 20, clip_04 20; nessuna nelle altre quattro.',
                          '')
    source=source.replace('- Test funzionali del codice: 12 superati nell\'ultima verifica documentata;\n  non sono risultati di accuratezza del sistema.','')
    source=source.replace('- Piano: 16/40 frame completi secondo il controllo strutturale. clip_01 completa,\n  inclusa la sequenza di otto frame; clip_03 e clip_04 complete per il piano spaziale.',
                          '- Piano: 40/40 frame completi secondo il controllo strutturale, incluse le tre sequenze di otto frame con le coppie richieste.')
    source=source.replace('Al momento ho 104 etichette,\n84 con coordinate annotabili, distribuite su tre clip; l\'annotazione è in corso.',
                          'Ho completato 216 etichette,\n176 con coordinate annotabili, distribuite sulle sette clip; resta la revisione qualitativa delle identità e del protocollo.')
    source=source.replace('Restano 24 frame del piano, nelle clip 05–08. Verificare inquadrature, movimento,',
                          'Il piano manuale è completato e non richiede altri punti. Verificare inquadrature, movimento,')
    source=source.replace('1. Completare il piano manuale e verificare coerenza delle due identità nelle sequenze.',
                          '1. Conservare la versione annotata completata e verificare coerenza delle due identità nelle sequenze.')
    source=source.replace('Eventuali estensioni: re-detection/re-prompt',
                          'Il notebook originale del tutorial contiene anche Clean player movement paths con clean_paths: può diventare un confronto aggiuntivo della correzione, a dati e protocollo fissi, prima di introdurre nuovi modelli. Questo confronto non è ancora implementato o valutato.\n\nEventuali estensioni: re-detection/re-prompt')
    source=source.replace('La pipeline e il correttore sono operativi sulla clip dimostrativa. Ho separato',
                          'La pipeline e il correttore sono operativi sulla clip dimostrativa e le annotazioni delle sette clip sono complete. Ho separato')
    source=source.replace('1. Completare il piano manuale','1. Rivedere il piano manuale')
    # Contenuto aggiornato consultabile senza Word; la versione iniziale resta intatta.
    updated=ROOT/'reports/GUIDA_INCONTRO_PROFESSORE_ANNOTAZIONI_COMPLETE_2026-10-20.md'
    updated.write_text(source,encoding='utf-8')
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.72)
    section.left_margin=section.right_margin=Inches(.85)
    styles=doc.styles
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
        styles[name].font.name='Calibri';styles[name].font.color.rgb=RGBColor(0,0,0)
    normal=styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.05;normal.paragraph_format.space_after=Pt(4)
    styles['Title'].font.size=Pt(23);styles['Title'].paragraph_format.space_after=Pt(10)
    for name,size in [('Heading 1',15),('Heading 2',12)]:
        styles[name].font.size=Pt(size);styles[name].font.bold=True
        styles[name].paragraph_format.space_before=Pt(11)
        styles[name].paragraph_format.space_after=Pt(5)
        styles[name].paragraph_format.keep_with_next=True
    doc.add_paragraph('Guida per il colloquio di tesi sul tracking nel basket','Title')
    doc.add_paragraph('Stato del progetto al 20 ottobre 2026','Subtitle')
    doc.add_paragraph('La pipeline e il correttore funzionano sulla clip dimostrativa. Le annotazioni delle sette clip sono complete; il prossimo passaggio è definire il protocollo e misurare il confronto grezzo e corretto sulle clip personali.')
    doc.core_properties.title='Guida per il colloquio di tesi sul tracking nel basket'
    doc.core_properties.subject='Stato verificato del progetto e scaletta notebook per notebook'
    doc.core_properties.author='Marco'
    lines=source.splitlines();i=1
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('## '):
            heading=line[3:]
            if heading.startswith('Notebook '):
                match=re.match(r'Notebook (\d+)\s+—\s+(.*)',heading)
                number,filename=match.groups()
                title={'01':'Pipeline di tracking e geometria','02':'Valutazione del detector','03':'Correttore e confronto sperimentale','04':'Dataset e annotazioni'}[number]
                doc.add_heading(f'Notebook {number} {title}',1)
                p=doc.add_paragraph(filename);p.paragraph_format.space_after=Pt(7)
                for run in p.runs:run.italic=True;run.font.size=Pt(10.5)
            else:
                doc.add_heading(re.sub(r'[^\w\sÀ-ÿ]','',heading),1)
            i+=1;continue
        if line.startswith('### '):
            doc.add_heading(re.sub(r'[^\w\sÀ-ÿ]','',line[4:]),2);i+=1;continue
        bullet=line.startswith('- ');numbered=bool(re.match(r'^\d+\. ',line))
        paragraph=[line[2:] if bullet else line]
        i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#{1,3} |\- |\d+\. )',lines[i].strip()):
            paragraph.append(lines[i].strip());i+=1
        text=' '.join(paragraph)
        p=doc.add_paragraph(style='List Bullet' if bullet else 'Normal')
        if numbered:
            p.paragraph_format.left_indent=Inches(.16)
            p.paragraph_format.first_line_indent=Inches(-.16)
        add_runs(p,text)
        if 'Split validation/test e categorie di inquadratura/movimento ancora vuoti.' in text:
            table(doc,stats)
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('Pagina ');run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    dest=ROOT/'exports/Incontro_professore_tracking_basket.docx'
    # Il template Word incorporato può avere una riga blu sotto il Title.
    for tree in [doc.styles.element,doc.element]:
        for border in tree.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    dest.parent.mkdir(exist_ok=True)
    doc.save(dest)
    print(dest)


if __name__=='__main__':main()
